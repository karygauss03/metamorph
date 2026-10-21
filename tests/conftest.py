"""Generate one small sample file per supported input format."""

import csv
import os
import subprocess
import zipfile
from pathlib import Path

import pytest

from metamorph import util


def _media(path, args):
    subprocess.run([util.ffmpeg(), "-hide_banner", "-loglevel", "error", "-y", *args, str(path)], check=True)


def make_fixtures(root):
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    from PIL import Image, ImageDraw
    import pillow_heif
    pillow_heif.register_heif_opener()

    img = Image.new("RGBA", (320, 200), (255, 255, 255, 0))
    d = ImageDraw.Draw(img)
    d.rectangle((40, 30, 280, 170), fill=(255, 138, 31, 255))
    d.ellipse((120, 60, 200, 140), fill=(30, 30, 30, 255))
    for ext, fmt in (("png", "PNG"), ("webp", "WEBP"), ("tiff", "TIFF"), ("gif", "GIF"), ("avif", "AVIF")):
        img.save(root / f"sample.{ext}", fmt)
    rgb = img.convert("RGB")
    rgb.save(root / "sample.jpg", "JPEG")
    rgb.save(root / "sample.bmp", "BMP")
    rgb.save(root / "sample.heic", "HEIF")
    (root / "sample.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="120" height="80">'
        '<rect width="120" height="80" fill="#ff8a1f"/><circle cx="60" cy="40" r="25" fill="#222"/></svg>')

    tone = ["-f", "lavfi", "-i", "sine=frequency=440:duration=2"]
    for ext in ("mp3", "m4a", "wav", "flac", "ogg", "opus", "aiff", "wma", "aac"):
        _media(root / f"sample.{ext}", tone)
    video = ["-f", "lavfi", "-i", "testsrc=size=320x240:rate=25:duration=2", *tone]
    for ext, extra in (("mp4", ["-pix_fmt", "yuv420p"]), ("mov", ["-pix_fmt", "yuv420p"]), ("mkv", []),
                       ("webm", ["-c:v", "libvpx-vp9", "-deadline", "realtime", "-cpu-used", "8"]),
                       ("avi", []), ("wmv", ["-c:v", "wmv2", "-c:a", "wmav2"]), ("mpg", []),
                       ("3gp", ["-c:v", "mpeg4", "-c:a", "aac", "-ar", "8000", "-ac", "1", "-b:a", "12k"]),
                       ("flv", ["-c:a", "libmp3lame", "-ar", "44100"])):
        _media(root / f"sample.{ext}", [*video, "-shortest", *extra])

    import pymupdf
    doc = pymupdf.open()
    for n in range(2):
        page = doc.new_page()
        page.insert_text((72, 80), f"Report page {n + 1}", fontsize=24)
        page.insert_text((72, 120), "Body text paragraph for the converter test.", fontsize=11)
        y = 160
        for r, row in enumerate([("Item", "Qty", "Price"), ("Apple", "3", "1.20"), ("Pear", "5", "0.80")]):
            for c, cell in enumerate(row):
                page.draw_rect(pymupdf.Rect(72 + c * 100, y + r * 20, 172 + c * 100, y + 20 + r * 20))
                page.insert_text((76 + c * 100, y + 14 + r * 20), cell, fontsize=10)
        page.insert_image(pymupdf.Rect(72, 300, 232, 400), filename=str(root / "sample.png"))
    doc.save(root / "sample.pdf")

    md = "# Title\n\nSome *emphasis* and a list:\n\n- one\n- two\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"
    (root / "sample.md").write_text(md)
    (root / "sample.txt").write_text("Plain text line one\nline two\n\n# not a heading\n")
    (root / "sample.html").write_text("<html><head><title>T</title></head><body><h1>Hello</h1><p>World</p></body></html>")
    (root / "sample.rtf").write_text(r"{\rtf1\ansi{\fonttbl\f0\fswiss Helvetica;}\f0\pard Hello {\b bold} RTF.\par}")
    pandoc = util.pandoc()
    subprocess.run([pandoc, str(root / "sample.md"), "-o", str(root / "sample.docx")], check=True)
    subprocess.run([pandoc, str(root / "sample.md"), "-o", str(root / "sample.odt")], check=True)
    subprocess.run([pandoc, str(root / "sample.md"), "-M", "title=Sample", "-o", str(root / "sample.epub")], check=True)

    from pptx import Presentation
    prs = Presentation()
    for title in ("First slide", "Second slide"):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        s.placeholders[1].text = "Bullet"
    prs.save(root / "sample.pptx")

    rows = [["name", "qty", "price"], ["apple", 3, 1.2], ["pear", 5, 0.8]]
    with open(root / "sample.csv", "w", newline="") as fh:
        csv.writer(fh).writerows(rows)
    with open(root / "sample.tsv", "w", newline="") as fh:
        csv.writer(fh, delimiter="\t").writerows(rows)
    from openpyxl import Workbook
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    wb.create_sheet("Second").append(["x", "y"])
    wb.save(root / "sample.xlsx")

    (root / "sample.srt").write_text("1\n00:00:01,000 --> 00:00:02,500\nHello <i>there</i>\n\n"
                                     "2\n00:00:03,000 --> 00:00:04,000\nSecond line\n")
    (root / "sample.vtt").write_text("WEBVTT\n\n00:01.000 --> 00:02.500\nHello\n\n00:03.000 --> 00:04.000\nAgain\n")

    with zipfile.ZipFile(root / "sample.zip", "w") as zf:
        zf.writestr("folder/a.txt", "a")
        zf.writestr("b.txt", "b")
    subprocess.run(["tar", "-cf", str(root / "sample.tar"), "-C", str(root), "sample.txt", "sample.md"], check=True)
    subprocess.run(["tar", "-czf", str(root / "sample.tar.gz"), "-C", str(root), "sample.txt"], check=True)
    subprocess.run(["gzip", "-kf", str(root / "sample.txt")], check=True)
    (root / "sample.txt.gz").rename(root / "notes.gz")

    # Office round trips need LibreOffice: build legacy formats from modern ones.
    for src, fmt, part in (("sample.docx", "doc", "writer"), ("sample.pptx", "ppt", "impress"),
                           ("sample.pptx", "odp", "impress"), ("sample.xlsx", "xls", "calc"),
                           ("sample.xlsx", "ods", "calc")):
        if util.soffice_component(part):
            util.soffice_convert(root / src, root, fmt)
    return root


@pytest.fixture(scope="session")
def fixtures(tmp_path_factory):
    # One LibreOffice profile per test worker: a profile can't be shared by two instances.
    os.environ["METAMORPH_LO_PROFILE"] = str(tmp_path_factory.mktemp("lo-profile"))
    return make_fixtures(tmp_path_factory.mktemp("fixtures"))
