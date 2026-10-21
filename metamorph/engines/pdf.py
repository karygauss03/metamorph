"""PDF conversion and tools (PyMuPDF, pdf2docx, python-pptx, openpyxl, Ghostscript)."""

import statistics
from pathlib import Path

from ..util import (ConversionError, ghostscript, parse_number, parse_page_ranges, run,
                    temp_dir, unique_path)
from . import documents

try:
    import pymupdf as fitz
except ImportError:  # older PyMuPDF
    import fitz

DPI = 300


def _open(src):
    try:
        doc = fitz.open(str(src))
    except Exception as exc:
        raise ConversionError(f"Cannot open PDF {Path(src).name}: {exc}")
    if doc.needs_pass:
        raise ConversionError(f"{Path(src).name} is password protected")
    return doc


def to_images(src, target, dest, ctx, dpi=DPI):
    """Render every page. One page -> a file; several -> a folder of files."""
    doc = _open(src)
    dest = Path(dest)
    if doc.page_count == 1:
        _render(doc[0], dest, target, dpi, ctx)
        return dest
    folder = unique_path(dest.with_suffix(""))
    folder.mkdir()
    width = len(str(doc.page_count))
    for i, page in enumerate(doc):
        _render(page, folder / f"{Path(src).stem}-{i + 1:0{width}d}.{target}", target, dpi, ctx)
        ctx.progress((i + 1) / doc.page_count)
    return folder


def _render(page, path, target, dpi, ctx):
    pix = page.get_pixmap(dpi=dpi, alpha=False)
    if target == "png":
        pix.save(str(path))
    else:
        pix.save(str(path), jpg_quality=ctx.settings.get("image_quality", 90))


def to_text(src, target, dest, ctx):
    doc = _open(src)
    text = "\n\f".join(page.get_text("text", sort=True) for page in doc)
    Path(dest).write_text(text, encoding="utf-8")
    return Path(dest)


def _markdown(doc):
    """Reconstruct Markdown with headings inferred from font size."""
    sizes = []
    for page in doc:
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    if span["text"].strip():
                        sizes.extend([round(span["size"], 1)] * len(span["text"]))
    body = statistics.mode(sizes) if sizes else 11
    out = []
    for page in doc:
        for block in page.get_text("dict", sort=True)["blocks"]:
            if block.get("type") != 0:
                continue
            lines = []
            block_size = 0
            for line in block["lines"]:
                parts = []
                for span in line["spans"]:
                    t = span["text"]
                    if not t.strip():
                        parts.append(t)
                        continue
                    block_size = max(block_size, span["size"])
                    bold = span["flags"] & 16 or "Bold" in span["font"]
                    parts.append(f"**{t.strip()}** " if bold and len(t.strip()) > 1 else t)
                lines.append("".join(parts).strip())
            text = " ".join(l for l in lines if l)
            if not text:
                continue
            ratio = block_size / body if body else 1
            if ratio >= 1.15 and len(text) < 200:
                level = 1 if ratio >= 1.8 else 2 if ratio >= 1.4 else 3
                out.append("#" * level + " " + text.replace("**", "").strip())
            elif text.lstrip().startswith(("•", "◦", "▪", "‣", "- ")):
                out.append("- " + text.lstrip("•◦▪‣- ").strip())
            else:
                out.append(text)
        out.append("")
    return "\n\n".join(s for s in out).strip() + "\n"


def to_markdown(src, target, dest, ctx):
    Path(dest).write_text(_markdown(_open(src)), encoding="utf-8")
    return Path(dest)


def to_docx(src, target, dest, ctx):
    try:
        from pdf2docx import Converter
    except ImportError:
        raise ConversionError("PDF to Word needs pdf2docx")
    import logging
    # pdf2docx logs progress through the root logger at INFO level.
    logging.disable(logging.INFO)
    try:
        conv = Converter(str(src))
        try:
            conv.convert(str(dest))
        finally:
            conv.close()
    finally:
        logging.disable(logging.NOTSET)
    return Path(dest)


def to_odt(src, target, dest, ctx):
    with temp_dir() as tmp:
        docx = to_docx(src, "docx", Path(tmp) / (Path(src).stem + ".docx"), ctx)
        return documents.pandoc_convert(docx, "odt", dest, ctx)


def to_epub(src, target, dest, ctx):
    with temp_dir() as tmp:
        try:
            docx = to_docx(src, "docx", Path(tmp) / (Path(src).stem + ".docx"), ctx)
            return documents.pandoc_convert(docx, "epub", dest, ctx, title=Path(src).stem)
        except ConversionError:
            md = Path(tmp) / (Path(src).stem + ".md")
            md.write_text(_markdown(_open(src)), encoding="utf-8")
            return documents.pandoc_convert(md, "epub", dest, ctx, title=Path(src).stem)


def to_pptx(src, target, dest, ctx):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        raise ConversionError("PDF to PowerPoint needs python-pptx")
    doc = _open(src)
    first = doc[0].rect
    prs = Presentation()
    # 1 pt = 12700 EMU; keep the PDF's aspect ratio.
    prs.slide_width, prs.slide_height = Emu(int(first.width * 12700)), Emu(int(first.height * 12700))
    blank = prs.slide_layouts[6]
    with temp_dir() as tmp:
        for i, page in enumerate(doc):
            img = Path(tmp) / f"p{i}.png"
            page.get_pixmap(dpi=200, alpha=False).save(str(img))
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(img), 0, 0, prs.slide_width, prs.slide_height)
            notes = page.get_text().strip()
            if notes:
                slide.notes_slide.notes_text_frame.text = notes
            ctx.progress((i + 1) / doc.page_count)
        prs.save(str(dest))
    return Path(dest)


def to_xlsx(src, target, dest, ctx):
    from openpyxl import Workbook
    doc = _open(src)
    wb = Workbook()
    wb.remove(wb.active)
    for i, page in enumerate(doc):
        rows = []
        try:
            for table in page.find_tables().tables:
                rows.extend(table.extract())
                rows.append([])
        except Exception:
            rows = []
        if not rows:
            # No ruled tables: fall back to word positions grouped into lines/columns.
            lines = {}
            for x0, y0, x1, y1, word, *_ in page.get_text("words", sort=True):
                lines.setdefault(round(y0 / 3), []).append((x0, x1, word))
            for _, words in sorted(lines.items()):
                cells, last_x1 = [], None
                for x0, x1, word in sorted(words):
                    if last_x1 is not None and x0 - last_x1 < 8:
                        cells[-1] += " " + word
                    else:
                        cells.append(word)
                    last_x1 = x1
                rows.append(cells)
        ws = wb.create_sheet(f"Page {i + 1}")
        for row in rows:
            ws.append([parse_number(c) for c in row])
    wb.save(str(dest))
    return Path(dest)



CONVERTERS = {
    "jpg": to_images, "png": to_images, "txt": to_text, "md": to_markdown, "docx": to_docx,
    "odt": to_odt, "epub": to_epub, "pptx": to_pptx, "xlsx": to_xlsx,
}


def convert(src, target, dest, ctx):
    return CONVERTERS[target](src, target, dest, ctx)


# ---------------------------------------------------------------- tools

def _dest(src, ctx, suffix, ext="pdf"):
    src = Path(src)
    return unique_path(Path(ctx.out_dir or src.parent) / f"{src.stem}{suffix}.{ext}")


def compress(src, ctx):
    dest = _dest(src, ctx, " (compressed)")
    level = {"Small": "/screen", "Balanced": "/ebook", "High": "/printer"}[ctx.params.get("level", "Balanced")]
    if ghostscript():
        run([ghostscript(), "-sDEVICE=pdfwrite", "-dCompatibilityLevel=1.6", f"-dPDFSETTINGS={level}",
             "-dNOPAUSE", "-dBATCH", "-dQUIET", "-dDetectDuplicateImages=true",
             f"-sOutputFile={dest}", str(src)])
    else:
        doc = _open(src)
        doc.rewrite_images(dpi_threshold=200, dpi_target=150, quality=70)
        doc.save(str(dest), garbage=4, deflate=True, clean=True)
    if dest.stat().st_size >= Path(src).stat().st_size:
        # Never hand back a "compressed" file that is larger.
        doc = _open(src)
        dest.unlink()
        doc.save(str(dest), garbage=4, deflate=True)
    return dest


def merge(paths, ctx):
    out = fitz.open()
    for i, p in enumerate(paths):
        out.insert_pdf(_open(p))
        ctx.progress((i + 1) / len(paths))
    dest = _dest(paths[0], ctx, " (merged)")
    out.save(str(dest), garbage=3, deflate=True)
    return dest


def split(src, ctx):
    doc = _open(src)
    src = Path(src)
    ranges = str(ctx.params.get("ranges", "")).strip()
    folder = unique_path(Path(ctx.out_dir or src.parent) / f"{src.stem} (split)")
    folder.mkdir()
    if ranges:
        groups = [parse_page_ranges(chunk, doc.page_count) for chunk in ranges.split(";") if chunk.strip()]
    else:
        groups = [[i] for i in range(doc.page_count)]
    width = len(str(doc.page_count))
    for n, pages in enumerate(groups):
        part = fitz.open()
        for p in pages:
            part.insert_pdf(doc, from_page=p, to_page=p)
        label = f"{pages[0] + 1:0{width}d}" if len(pages) == 1 else f"{pages[0] + 1}-{pages[-1] + 1}"
        part.save(str(folder / f"{src.stem} p{label}.pdf"), garbage=3, deflate=True)
        ctx.progress((n + 1) / len(groups))
    return folder


def extract_pages(src, ctx):
    doc = _open(src)
    pages = parse_page_ranges(ctx.params.get("pages", "1"), doc.page_count)
    out = fitz.open()
    for p in pages:
        out.insert_pdf(doc, from_page=p, to_page=p)
    dest = _dest(src, ctx, " (pages)")
    out.save(str(dest), garbage=3, deflate=True)
    return dest


def remove_pages(src, ctx):
    doc = _open(src)
    drop = set(parse_page_ranges(ctx.params.get("pages", "1"), doc.page_count))
    if len(drop) == doc.page_count:
        raise ConversionError("Cannot remove every page")
    doc.select([i for i in range(doc.page_count) if i not in drop])
    dest = _dest(src, ctx, " (edited)")
    doc.save(str(dest), garbage=3, deflate=True)
    return dest


def rotate(src, ctx):
    doc = _open(src)
    angle = {"90° clockwise": 90, "90° counter-clockwise": 270, "180°": 180}[ctx.params.get("angle", "90° clockwise")]
    for page in doc:
        page.set_rotation((page.rotation + angle) % 360)
    dest = _dest(src, ctx, " (rotated)")
    doc.save(str(dest), garbage=3, deflate=True)
    return dest


def extract_images(src, ctx):
    doc = _open(src)
    src = Path(src)
    folder = unique_path(Path(ctx.out_dir or src.parent) / f"{src.stem} (images)")
    folder.mkdir()
    seen, count = set(), 0
    for page_no, page in enumerate(doc):
        for img in page.get_images(full=True):
            xref = img[0]
            if xref in seen:
                continue
            seen.add(xref)
            data = doc.extract_image(xref)
            if not data or data["width"] < 16 or data["height"] < 16:
                continue
            count += 1
            (folder / f"page{page_no + 1}-{count}.{data['ext']}").write_bytes(data["image"])
    if not count:
        folder.rmdir()
        raise ConversionError(f"{src.name} contains no embedded images")
    return folder


def page_count(src):
    return _open(src).page_count
