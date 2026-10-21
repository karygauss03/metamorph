"""Presentations and spreadsheets (LibreOffice, openpyxl, csv)."""

import csv
import shutil
from pathlib import Path

from ..util import ConversionError, parse_number, soffice_convert, temp_dir, unique_path
from . import pdf

# LibreOffice CSV filter options: separator, quote char, UTF-8, first line.
CSV_FILTER = "Text - txt - csv (StarCalc):44,34,76,1"
TSV_FILTER = "Text - txt - csv (StarCalc):9,34,76,1"


def slides_convert(src, target, dest, ctx):
    src, dest = Path(src), Path(dest)
    with temp_dir() as tmp:
        if target in ("pdf", "pptx", "odp"):
            shutil.move(str(soffice_convert(src, tmp, target)), dest)
            return dest
        # Images: render via PDF so every slide becomes a picture.
        rendered = soffice_convert(src, tmp, "pdf")
        return pdf.to_images(rendered, target, dest, ctx, dpi=150)


def _read_rows(src):
    """Yield (sheet name, rows) for every sheet of a spreadsheet file."""
    src = Path(src)
    ext = src.suffix.lower().lstrip(".")
    if ext in ("csv", "tsv"):
        with open(src, newline="", encoding="utf-8-sig", errors="replace") as fh:
            sample = fh.read(8192)
            fh.seek(0)
            delim = "\t" if ext == "tsv" else ","
            if ext == "csv":
                try:
                    delim = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
                except csv.Error:
                    pass
            return [(src.stem, list(csv.reader(fh, delimiter=delim)))]
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(src, read_only=True, data_only=True)
        return [(ws.title, [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)])
                for ws in wb.worksheets]
    with temp_dir() as tmp:  # xls / ods: normalise through LibreOffice first
        return _read_rows(soffice_convert(src, tmp, "xlsx"))


def _write_delimited(sheets, target, dest):
    delim = "\t" if target == "tsv" else ","
    dest = Path(dest)
    if len(sheets) == 1:
        with open(dest, "w", newline="", encoding="utf-8") as fh:
            csv.writer(fh, delimiter=delim).writerows(sheets[0][1])
        return dest
    folder = unique_path(dest.with_suffix(""))
    folder.mkdir()
    for name, rows in sheets:
        safe = "".join(c if c.isalnum() or c in " -_" else "_" for c in str(name)).strip() or "Sheet"
        with open(folder / f"{safe}.{target}", "w", newline="", encoding="utf-8") as fh:
            csv.writer(fh, delimiter=delim).writerows(rows)
    return folder


def _typed(value):
    return parse_number(value) if isinstance(value, str) else value


def sheets_convert(src, target, dest, ctx):
    src, dest = Path(src), Path(dest)
    ext = src.suffix.lower().lstrip(".")
    if target in ("csv", "tsv"):
        return _write_delimited(_read_rows(src), target, dest)
    with temp_dir() as tmp:
        if ext in ("csv", "tsv"):
            if target == "xlsx":
                from openpyxl import Workbook
                wb = Workbook()
                ws = wb.active
                ws.title = src.stem[:31] or "Sheet1"
                for row in _read_rows(src)[0][1]:
                    ws.append([_typed(v) for v in row])
                wb.save(dest)
                return dest
            infilter = TSV_FILTER if ext == "tsv" else CSV_FILTER
            shutil.move(str(soffice_convert(src, tmp, target, infilter=infilter)), dest)
            return dest
        shutil.move(str(soffice_convert(src, tmp, target)), dest)
        return dest
